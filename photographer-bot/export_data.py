import json
import subprocess
from pathlib import Path
from urllib.parse import urlencode

ROOT = Path(__file__).resolve().parent.parent
BOT_DIR = Path(__file__).resolve().parent
DATA_DIR = BOT_DIR / "data"
PHOTOGRAPHERS_HTML = ROOT / "photographers-table.html"
COURSE_HTML = ROOT / "portrait-course.html"
WORKBOOK_PATH = ROOT / "photographers-and-course.xlsx"


def extract_js_value(path, start_marker, end_marker, expression):
    script = r"""
const fs = require('fs');
const vm = require('vm');
const [path, startMarker, endMarker, expression] = process.argv.slice(1);
const html = fs.readFileSync(path, 'utf8');
const start = html.indexOf(startMarker);
const end = html.indexOf(endMarker, start);
if (start < 0 || end < 0) throw new Error('Data markers not found');
const code = html.slice(start, end) + '\nJSON.stringify(' + expression + ')';
process.stdout.write(vm.runInNewContext(code));
"""
    result = subprocess.run(
        ["node", "-e", script, str(path), start_marker, end_marker, expression],
        check=True,
        capture_output=True,
        text=True,
    )
    return json.loads(result.stdout)


def photographer_rows():
    rows = extract_js_value(
        PHOTOGRAPHERS_HTML,
        "const masters=",
        "const labels=",
        "rows",
    )
    for row in rows:
        query = (row.get("latin") or row["photographer"]) + " фотограф официальный сайт"
        row["search_url"] = "https://www.google.com/search?" + urlencode({"q": query})
    return rows


def assignment_rows():
    topics = extract_js_value(COURSE_HTML, "const topics=", "const svgByPart=", "topics")
    course_url = COURSE_HTML.as_uri()
    lessons = []
    for topic_index, topic in enumerate(topics):
        lessons.append(
            {
                "part": topic["part"],
                "title": topic["title"],
                "hours": topic["hours"],
                "theory": topic["theory"],
                "practice": topic["practice"],
                "result": topic["result"],
                "advice": topic["advice"],
                "source": topic["source"],
                "topic_index": topic_index,
                "course_url": f"{course_url}#t{topic_index}",
            }
        )
    return lessons


def save_json(name, data):
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    (DATA_DIR / name).write_text(
        json.dumps(data, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def add_sheet(workbook, name, rows, columns, widths):
    from openpyxl.styles import Alignment, Font, PatternFill

    sheet = workbook.create_sheet(name)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{chr(64 + len(columns))}{len(rows) + 1}"
    header_fill = PatternFill("solid", fgColor="243149")
    for column_index, (key, label) in enumerate(columns, start=1):
        cell = sheet.cell(1, column_index, label)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[cell.column_letter].width = widths[column_index - 1]
    sheet.row_dimensions[1].height = 28
    for row_index, row in enumerate(rows, start=2):
        for column_index, (key, _) in enumerate(columns, start=1):
            value = row.get(key, "")
            if isinstance(value, list):
                value = "\n".join(
                    f"{index}. {item}" for index, item in enumerate(value, start=1)
                )
            cell = sheet.cell(row_index, column_index, value)
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if row_index % 2 == 0:
            for cell in sheet[row_index]:
                cell.fill = PatternFill("solid", fgColor="F3F6FA")
    return sheet


def save_workbook(photographers, assignments):
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    add_sheet(
        workbook,
        "Фотографы",
        photographers,
        [
            ("photographer", "Фотограф"),
            ("latin", "Имя латиницей"),
            ("section", "Раздел"),
            ("topic", "Тема"),
            ("style", "Стиль"),
            ("work", "Работа или чем известен"),
            ("source", "Источник"),
            ("search_url", "Поиск работ"),
        ],
        [24, 24, 18, 28, 28, 54, 48, 44],
    )
    add_sheet(
        workbook,
        "Задания курса",
        assignments,
        [
            ("part", "Раздел курса"),
            ("title", "Тема"),
            ("hours", "Часы"),
            ("theory", "Теория"),
            ("practice", "Практика"),
            ("result", "Цель"),
            ("advice", "Советы"),
            ("source", "Источник"),
            ("topic_index", "№ темы (с 0)"),
            ("course_url", "Ссылка на тему"),
        ],
        [24, 32, 10, 64, 64, 54, 54, 48, 16, 44],
    )
    workbook.save(WORKBOOK_PATH)


def main():
    photographers = photographer_rows()
    assignments = assignment_rows()
    save_json("photographers.json", photographers)
    save_json("assignments.json", assignments)
    save_workbook(photographers, assignments)
    print(
        json.dumps(
            {
                "photographers": len(photographers),
                "assignments": len(assignments),
                "workbook": str(WORKBOOK_PATH),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
